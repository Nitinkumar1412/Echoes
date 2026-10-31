# -*- coding: utf-8 -*-
"""Unit tests for document parsers (Text, PDF, DOCX, Factory)."""
import tempfile
from pathlib import Path

import docx
import pymupdf
import pytest

from backend.app.models.document import FileType
from backend.app.services.parsers import (
    detect_file_type,
    parse_docx_file,
    parse_document,
    parse_pdf_file,
    parse_text_file,
)


def test_text_parser():
    """Verify plain text and markdown parsing."""
    with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False, encoding="utf-8") as f:
        f.write("# Echoes Project\n\nThis is a personal knowledge retrieval engine.\nIt connects your past.")
        temp_path = f.name

    try:
        parsed = parse_text_file(temp_path)
        assert parsed.extracted_title == "Echoes Project"
        assert "personal knowledge retrieval engine" in parsed.text
        assert parsed.page_count == 1
        assert parsed.word_count > 0
        assert len(parsed.pages) == 1
    finally:
        Path(temp_path).unlink(missing_ok=True)


def test_pdf_parser():
    """Verify PyMuPDF digital PDF parsing with page tracking."""
    # Create a 2-page PDF in memory using pymupdf
    pdf_doc = pymupdf.open()
    
    page1 = pdf_doc.new_page()
    page1.insert_text((50, 72), "Page 1: Project Kickoff Meeting Notes.\nAttendees: Alice, Bob.")
    
    page2 = pdf_doc.new_page()
    page2.insert_text((50, 72), "Page 2: Decision made to use MongoDB Atlas and ChromaDB.")
    
    temp_f = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
    pdf_path = temp_f.name
    temp_f.close()
    pdf_doc.save(pdf_path)
    pdf_doc.close()

    try:
        parsed = parse_pdf_file(pdf_path)
        assert parsed.page_count == 2
        assert len(parsed.pages) == 2
        assert "Project Kickoff" in parsed.pages[0].text
        assert "ChromaDB" in parsed.pages[1].text
        assert parsed.pages[0].page_number == 1
        assert parsed.pages[1].page_number == 2
    finally:
        Path(pdf_path).unlink(missing_ok=True)


def test_docx_parser():
    """Verify Word .docx parsing with paragraphs and tables."""
    doc = docx.Document()
    doc.add_heading("Design Specification", level=1)
    doc.add_paragraph("This document outlines the architecture for Echoes Search.")
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Module"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Database"
    table.cell(1, 1).text = "Done"

    temp_docx = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    docx_path = temp_docx.name
    temp_docx.close()
    doc.save(docx_path)

    try:
        parsed = parse_docx_file(docx_path)
        assert "Design Specification" in parsed.text
        assert "architecture for Echoes Search" in parsed.text
        assert "Database | Done" in parsed.text
    finally:
        Path(docx_path).unlink(missing_ok=True)


def test_parser_factory():
    """Verify file type detection and factory routing."""
    assert detect_file_type("report.pdf") == FileType.PDF
    assert detect_file_type("notes.docx") == FileType.DOCX
    assert detect_file_type("readme.md") == FileType.TXT
    assert detect_file_type("screenshot.png") == FileType.IMAGE
    assert detect_file_type("unknown.xyz") == FileType.OTHER
