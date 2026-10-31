# -*- coding: utf-8 -*-
"""DOCX parser using python-docx to extract paragraphs, headings, and tables."""
from __future__ import annotations

from pathlib import Path
from typing import Union

import docx

from .base import PageContent, ParsedDocument


def parse_docx_file(file_path: Union[str, Path]) -> ParsedDocument:
    """Parse Microsoft Word .docx file into structured text."""
    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"DOCX file not found: {path}")

    doc = docx.Document(str(path))
    text_blocks: list[str] = []

    # 1. Paragraphs (with headings)
    for p in doc.paragraphs:
        p_text = p.text.strip()
        if not p_text:
            continue
        style_name = p.style.name.lower() if p.style and p.style.name else ""
        if "heading" in style_name:
            text_blocks.append(f"\n## {p_text}\n")
        else:
            text_blocks.append(p_text)

    # 2. Tables
    for table in doc.tables:
        table_rows: list[str] = []
        for row in table.rows:
            cell_texts = [c.text.strip() for c in row.cells if c.text.strip()]
            if cell_texts:
                table_rows.append(" | ".join(cell_texts))
        if table_rows:
            text_blocks.append("\n" + "\n".join(table_rows) + "\n")

    full_text = "\n\n".join(text_blocks).strip()

    # Core properties
    title = None
    try:
        core_props = doc.core_properties
        title = core_props.title if core_props and core_props.title else None
    except Exception:
        pass

    pages = [
        PageContent(
            page_number=1,
            text=full_text,
            char_start=0,
            char_end=len(full_text),
        )
    ]

    return ParsedDocument(
        text=full_text,
        pages=pages,
        page_count=1,
        word_count=len(full_text.split()),
        extracted_title=title or path.stem,
        metadata={"filename": path.name},
    )
